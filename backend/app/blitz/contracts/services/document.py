"""Договор и соглашение о конфиденциальности из шаблонов Word. У аудита свой договор и префикс номера.

В шаблонах вместо пропусков стоят метки вида {{customer_inn}}, каждая целиком внутри одного
фрагмента текста — замена не ломает оформление. Оба документа ссылаются на один номер договора.
"""

from datetime import datetime
from decimal import Decimal
from enum import StrEnum
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as WordDocument
from docx.text.paragraph import Paragraph

from app.blitz.catalog import CONTRACT_SUBJECTS, DEADLINE_DAYS, ContractKind, Deadline
from app.blitz.orders.models import COMPANY_CUSTOMERS, BlitzOrder, OrderCompany, OrderIndividual, ServiceKind
from app.core.money import amount_in_words, format_amount, number_to_words, plural, vat_included

TEMPLATES = Path(__file__).resolve().parents[1] / "templates"

MONTHS = (
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)
WORKING_DAYS = ("рабочий день", "рабочих дня", "рабочих дней")
NDA_PARTY = "Сторона, передающая информацию (Сторона 1), с одной стороны и"


class Paper(StrEnum):
    """Какой документ: договор или соглашение о конфиденциальности."""

    CONTRACT = "contract"
    NDA = "nda"


PAPER_TITLES = {Paper.CONTRACT: "Договор", Paper.NDA: "Соглашение о конфиденциальности"}


def contract_problem(order: BlitzOrder) -> str | None:
    """Чего не хватает для договора. None — всё на месте."""

    if order.customer_type in COMPANY_CUSTOMERS and order.company is None:
        return "В заявке нет реквизитов заказчика"

    if order.customer_type not in COMPANY_CUSTOMERS and order.individual is None:
        return "В заявке нет паспортных данных заказчика"

    if order.contract_kind is None:
        return "Вид договора ещё не определён"

    if order.price is None:
        return "Цена ещё не согласована"

    return None


def contract_number(order: BlitzOrder, signed_at: datetime) -> str:
    """Номер договора: БЭ — экспертиза, БА — аудит, затем год подписания и номер заявки."""

    prefix = "БА" if order.service == ServiceKind.AUDIT else "БЭ"

    return f"{prefix}-{signed_at.year}-{order.id:04d}"


def paper_name(paper: Paper, order: BlitzOrder, signed_at: datetime) -> str:
    """Имя файла: «Договор БЭ-2026-0001.docx»."""

    return f"{PAPER_TITLES[paper]} {contract_number(order, signed_at)}.docx"


def format_date(moment: datetime) -> str:
    """Дата как в договоре: «25» сентября 2026 г."""

    return f"«{moment.day:02d}» {MONTHS[moment.month - 1]} {moment.year} г."


def initials(full_name: str) -> str:
    """Иванов Иван Иванович → И.И. Иванов."""

    surname, *names = full_name.split()
    letters = "".join(f"{name[0]}." for name in names)

    return f"{letters} {surname}" if letters else surname


def working_days(deadline: str) -> str:
    """Срок из пожелания заказчика: «3 (три) рабочих дня»."""

    days = DEADLINE_DAYS[Deadline(deadline)]

    return f"{days} ({number_to_words(days)}) {plural(days, WORKING_DAYS)}"


def money_text(amount: Decimal) -> str:
    """Сумма цифрами и прописью: «20 000,00 (Двадцать тысяч рублей 00 копеек)»."""

    return f"{format_amount(amount)} ({amount_in_words(amount)})"


def vat_clause(amount: Decimal, rate: int) -> str:
    """Фраза про НДС в пункте о стоимости."""

    if rate <= 0:
        return "НДС не облагается"

    return f"в т.ч. НДС {rate}% – {money_text(vat_included(amount, rate))}"


def numbered(prefix: str, lines: list[str]) -> dict[str, str]:
    """Строки блока реквизитов как метки prefix_1, prefix_2…"""

    return {f"{prefix}_{index}": line for index, line in enumerate(lines, start=1)}


def company_values(company: OrderCompany, phone: str, email: str) -> dict[str, str]:
    """Преамбулы, реквизиты и подпись организации."""

    inn_line = f"ИНН/КПП {company.inn}/{company.kpp or '—'}"
    person = (
        f"{company.full_name} ({company.name}), в лице {company.signer_genitive}, "
        f"действующего на основании {company.signer_basis}"
    )
    contract_lines = [
        company.name,
        inn_line,
        f"ОГРН {company.ogrn}",
        f"Юридический/фактический адрес: {company.address}",
        f"р/с {company.account}",
        f"БАНК {company.bank}",
        f"к/с {company.corr_account}",
        f"БИК {company.bic}",
    ]
    nda_lines = [
        f"Наименование: {company.name}",
        f"Адрес, указанный в ЕГРЮЛ: {company.address}",
        "",
        f"Телефон: {phone}",
        f"Электронная почта: {email}",
        f"ОГРН {company.ogrn}",
        inn_line,
    ]

    return {
        "customer_preamble": f"{person}, именуемое в дальнейшем «Заказчик», с одной стороны,",
        "nda_preamble": f"{person}, далее именуемое {NDA_PARTY}",
        "customer_name": company.name,
        "signer_position": company.signer_position,
        "signer_initials": initials(company.signer_name),
        **numbered("customer_line", contract_lines),
        **numbered("nda_line", nda_lines),
    }


def individual_values(individual: OrderIndividual, phone: str, email: str) -> dict[str, str]:
    """Преамбулы, реквизиты и подпись физлица."""

    issued = f"{individual.passport_issued_by}, {individual.passport_issued_at:%d.%m.%Y}"
    person = (
        f"{individual.full_name}, паспорт {individual.passport_number}, выдан {issued}, "
        f"адрес регистрации: {individual.address}"
    )
    contract_lines = [
        individual.full_name,
        f"Паспорт: {individual.passport_number}",
        f"Выдан: {issued}",
        f"Адрес регистрации: {individual.address}",
        f"Телефон: {phone}",
        f"E-mail: {email}",
        "",
        "",
    ]
    nda_lines = [
        f"ФИО: {individual.full_name}",
        f"Адрес регистрации: {individual.address}",
        f"Паспорт: {individual.passport_number}, выдан {issued}",
        f"Телефон: {phone}",
        f"Электронная почта: {email}",
        "",
        "",
    ]

    return {
        "customer_preamble": f"{person}, именуемый в дальнейшем «Заказчик», с одной стороны,",
        "nda_preamble": f"{person}, далее именуемый {NDA_PARTY}",
        "customer_name": individual.full_name,
        "signer_position": "Заказчик",
        "signer_initials": initials(individual.full_name),
        **numbered("customer_line", contract_lines),
        **numbered("nda_line", nda_lines),
    }


def paper_values(order: BlitzOrder, vat_rate: int, signed_at: datetime) -> dict[str, str]:
    """Значения меток. Перед вызовом заявку проверяет contract_problem."""

    price = order.price or Decimal(0)
    phone = order.customer.phone or "—"
    email = order.customer.email or "—"

    if order.company is not None:
        party = company_values(order.company, phone, email)
    elif order.individual is not None:
        party = individual_values(order.individual, phone, email)
    else:
        raise ValueError("В заявке нет данных заказчика")

    signed_on = format_date(signed_at)

    return {
        "number": contract_number(order, signed_at),
        "date": signed_on,
        "date_caps": signed_on.upper(),
        "subject": CONTRACT_SUBJECTS[ContractKind(order.contract_kind)],
        "object_name": order.object_name,
        "days": working_days(order.deadline),
        "price": money_text(price),
        "vat": vat_clause(price, vat_rate),
        **party,
    }


def all_paragraphs(document: WordDocument) -> list[Paragraph]:
    """Абзацы документа вместе с абзацами в ячейках таблиц."""

    cells = [
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]

    return [*document.paragraphs, *cells]


def fill(document: WordDocument, values: dict[str, str]) -> None:
    """Подставить значения вместо меток."""

    for paragraph in all_paragraphs(document):
        for run in paragraph.runs:
            if "{{" not in run.text:
                continue

            text = run.text
            for key, value in values.items():
                text = text.replace("{{" + key + "}}", value)
            run.text = text


def template_of(paper: Paper, order: BlitzOrder) -> Path:
    """Шаблон документа: у договора аудита свой."""

    if paper == Paper.CONTRACT and order.service == ServiceKind.AUDIT:
        return TEMPLATES / "contract_audit.docx"

    return TEMPLATES / f"{paper}.docx"


def build_paper(paper: Paper, order: BlitzOrder, vat_rate: int, signed_at: datetime) -> bytes:
    """Собрать договор или соглашение о конфиденциальности."""

    document = Document(str(template_of(paper, order)))
    fill(document, paper_values(order, vat_rate, signed_at))

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()
