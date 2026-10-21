"""Отчёт о представленных документах: какие пункты перечня закрыты файлами, а какие нет."""

from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

from app.blitz.audit.checklist import AUDIT_DOCUMENTS
from app.blitz.documents.models import DocumentKind, OrderDocument
from app.blitz.orders.models import BlitzOrder

PROVIDED = "Представлен"
MISSING = "Не представлен"
COLUMN_WIDTHS = (Cm(1), Cm(13), Cm(3), Cm(8))


def files_by_item(documents: list[OrderDocument]) -> dict[int, list[str]]:
    """Имена файлов по номерам пунктов перечня."""

    grouped: dict[int, list[str]] = {}

    for document in documents:
        if document.kind == DocumentKind.AUDIT_ITEM and document.item_number is not None:
            grouped.setdefault(document.item_number, []).append(document.name)

    return grouped


def build_documents_report(order: BlitzOrder, generated_at: datetime) -> bytes:
    """Word-файл: шапка, итог по комплектности и таблица по всем пунктам перечня."""

    grouped = files_by_item(order.documents)
    total = len(AUDIT_DOCUMENTS)
    document = Document()

    section = document.sections[0]
    width, height = section.page_width, section.page_height
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = height, width
    section.left_margin, section.right_margin = Cm(2), Cm(1.5)

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    document.add_heading("Отчёт о представленных документах", level=1)
    document.add_paragraph("для аудита системы управления промышленной безопасностью")
    document.add_paragraph(f"Заявка № {order.id}")
    customer = order.company.name if order.company else order.customer.display_name
    document.add_paragraph(f"Заказчик: {customer}")
    document.add_paragraph(f"Объект аудита: {order.object_name}")
    document.add_paragraph(f"Дата формирования: {generated_at:%d.%m.%Y %H:%M}")

    summary = document.add_paragraph()
    summary.add_run(
        f"Представлены документы по {len(grouped)} из {total} пунктов перечня, "
        f"не представлены — по {total - len(grouped)}.",
    ).bold = True

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    for cell, title in zip(table.rows[0].cells, ("№", "Документ по перечню", "Статус", "Файлы"), strict=True):
        cell.text = title

    for number, title in enumerate(AUDIT_DOCUMENTS, start=1):
        names = grouped.get(number, [])
        row = table.add_row().cells
        row[0].text = str(number)
        row[1].text = title
        row[2].text = PROVIDED if names else MISSING
        row[3].text = "\n".join(names)

    for table_row in table.rows:
        for cell, column_width in zip(table_row.cells, COLUMN_WIDTHS, strict=True):
            cell.width = column_width

    revisions = [item for item in order.documents if item.kind == DocumentKind.REVISION]
    if revisions:
        document.add_heading("Документы, присланные после замечаний", level=2)
        for item in revisions:
            document.add_paragraph(item.name, style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def report_filename(order: BlitzOrder) -> str:
    """Имя файла отчёта."""

    return f"Отчёт о представленных документах — заявка {order.id}.docx"
